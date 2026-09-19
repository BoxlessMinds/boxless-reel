"""File processing utilities for document text extraction."""

import logging
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO

logger = logging.getLogger(__name__)

# Supported file types and their MIME types
SUPPORTED_FILE_TYPES = {
    "pdf": ["application/pdf"],
    "docx": ["application/vnd.openxmlformats-officedocument.wordprocessingml.document"],
    "txt": ["text/plain"],
    "md": ["text/markdown", "text/plain"],
}

# File size limits (in bytes)
FILE_SIZE_LIMITS = {
    "pdf": 10 * 1024 * 1024,  # 10MB
    "docx": 10 * 1024 * 1024,  # 10MB
    "txt": 5 * 1024 * 1024,  # 5MB
    "md": 5 * 1024 * 1024,  # 5MB
}


class FileProcessingError(Exception):
    """Base exception for file processing errors."""

    pass


class UnsupportedFileTypeError(FileProcessingError):
    """Raised when the file type is not supported."""

    pass


class FileTooLargeError(FileProcessingError):
    """Raised when the file exceeds size limits."""

    pass


class ExtractionError(FileProcessingError):
    """Raised when text extraction fails."""

    pass


@dataclass
class ExtractedPage:
    """A single page of extracted content."""

    page_number: int
    text: str


@dataclass
class ExtractedContent:
    """Result of text extraction from a document."""

    text: str
    page_count: int | None
    pages: list[ExtractedPage] | None
    word_count: int

    @classmethod
    def from_text(cls, text: str) -> "ExtractedContent":
        """Create ExtractedContent from plain text."""
        word_count = len(text.split())
        return cls(
            text=text,
            page_count=None,
            pages=None,
            word_count=word_count,
        )

    @classmethod
    def from_pages(cls, pages: list[ExtractedPage]) -> "ExtractedContent":
        """Create ExtractedContent from a list of pages."""
        full_text = "\n\n".join(page.text for page in pages)
        word_count = len(full_text.split())
        return cls(
            text=full_text,
            page_count=len(pages),
            pages=pages,
            word_count=word_count,
        )


class FileProcessor:
    """Processor for extracting text from various document formats."""

    def __init__(self) -> None:
        """Initialize the file processor."""
        pass

    def get_file_type(self, filename: str) -> str | None:
        """
        Get the file type from filename extension.

        Args:
            filename: The filename with extension.

        Returns:
            File type string (pdf, docx, txt, md) or None if unsupported.
        """
        ext = Path(filename).suffix.lower().lstrip(".")
        if ext in SUPPORTED_FILE_TYPES:
            return ext
        return None

    def validate_file_type(self, filename: str) -> str:
        """
        Validate that the file type is supported.

        Args:
            filename: The filename with extension.

        Returns:
            The validated file type.

        Raises:
            UnsupportedFileTypeError: If file type is not supported.
        """
        file_type = self.get_file_type(filename)
        if file_type is None:
            supported = ", ".join(SUPPORTED_FILE_TYPES.keys())
            raise UnsupportedFileTypeError(
                f"Unsupported file type. Supported types: {supported}"
            )
        return file_type

    # Magic byte signatures for content validation
    MAGIC_BYTES = {
        "pdf": b"%PDF",
        "docx": b"PK",  # DOCX is a ZIP archive
    }

    def validate_file_content(self, content: bytes, file_type: str) -> bool:
        """
        Validate that file content matches the expected type via magic bytes.

        Args:
            content: Raw file bytes (at least first 4 bytes needed).
            file_type: Expected file type (pdf, docx, txt, md).

        Returns:
            True if content matches expected type.
        """
        expected = self.MAGIC_BYTES.get(file_type)
        if expected:
            return content[: len(expected)] == expected
        if file_type in ("txt", "md"):
            try:
                content[:512].decode("utf-8")
                return True
            except UnicodeDecodeError:
                return False
        return True

    def validate_file_size(self, file_type: str, size_bytes: int) -> None:
        """
        Validate that the file size is within limits.

        Args:
            file_type: The file type (pdf, docx, txt, md).
            size_bytes: File size in bytes.

        Raises:
            FileTooLargeError: If file exceeds size limit.
        """
        limit = FILE_SIZE_LIMITS.get(file_type, 5 * 1024 * 1024)
        if size_bytes > limit:
            limit_mb = limit / (1024 * 1024)
            raise FileTooLargeError(
                f"File too large. Maximum size for {file_type} files is {limit_mb:.0f}MB"
            )

    def extract_text(self, file_path: Path, file_type: str) -> ExtractedContent:
        """
        Extract text content from a document file.

        Args:
            file_path: Path to the file.
            file_type: Type of file (pdf, docx, txt, md).

        Returns:
            ExtractedContent with text and metadata.

        Raises:
            ExtractionError: If text extraction fails.
        """
        logger.debug("Extracting text from %s (type=%s)", file_path, file_type)

        try:
            if file_type == "pdf":
                return self._extract_pdf(file_path)
            elif file_type == "docx":
                return self._extract_docx(file_path)
            elif file_type in ("txt", "md"):
                return self._extract_text_file(file_path)
            else:
                raise UnsupportedFileTypeError(f"Unsupported file type: {file_type}")
        except (UnsupportedFileTypeError, ExtractionError):
            raise
        except Exception as e:
            logger.error("Failed to extract text from %s: %s", file_path, e)
            raise ExtractionError(f"Failed to extract text: {e}") from e

    def extract_text_from_stream(
        self, file_stream: BinaryIO, file_type: str
    ) -> ExtractedContent:
        """
        Extract text content from a file stream.

        Args:
            file_stream: Binary file stream.
            file_type: Type of file (pdf, docx, txt, md).

        Returns:
            ExtractedContent with text and metadata.

        Raises:
            ExtractionError: If text extraction fails.
        """
        try:
            if file_type == "pdf":
                return self._extract_pdf_from_stream(file_stream)
            elif file_type == "docx":
                return self._extract_docx_from_stream(file_stream)
            elif file_type in ("txt", "md"):
                return self._extract_text_from_stream(file_stream)
            else:
                raise UnsupportedFileTypeError(f"Unsupported file type: {file_type}")
        except (UnsupportedFileTypeError, ExtractionError):
            raise
        except Exception as e:
            logger.error("Failed to extract text from stream: %s", e)
            raise ExtractionError(f"Failed to extract text: {e}") from e

    def _extract_pdf(self, file_path: Path) -> ExtractedContent:
        """Extract text from a PDF file."""
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ExtractionError(
                "pypdf is required for PDF processing. Install with: uv add pypdf"
            )

        reader = PdfReader(file_path)
        pages = []

        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(ExtractedPage(page_number=i, text=text.strip()))

        if not pages:
            raise ExtractionError("No text could be extracted from PDF")

        return ExtractedContent.from_pages(pages)

    def _extract_pdf_from_stream(self, file_stream: BinaryIO) -> ExtractedContent:
        """Extract text from a PDF file stream."""
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ExtractionError(
                "pypdf is required for PDF processing. Install with: uv add pypdf"
            )

        reader = PdfReader(file_stream)
        pages = []

        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(ExtractedPage(page_number=i, text=text.strip()))

        if not pages:
            raise ExtractionError("No text could be extracted from PDF")

        return ExtractedContent.from_pages(pages)

    def _extract_docx(self, file_path: Path) -> ExtractedContent:
        """Extract text from a DOCX file."""
        try:
            from docx import Document
        except ImportError:
            raise ExtractionError(
                "python-docx is required for DOCX processing. Install with: uv add python-docx"
            )

        doc = Document(file_path)
        paragraphs = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        if not paragraphs:
            raise ExtractionError("No text could be extracted from DOCX")

        full_text = "\n\n".join(paragraphs)
        return ExtractedContent.from_text(full_text)

    def _extract_docx_from_stream(self, file_stream: BinaryIO) -> ExtractedContent:
        """Extract text from a DOCX file stream."""
        try:
            from docx import Document
        except ImportError:
            raise ExtractionError(
                "python-docx is required for DOCX processing. Install with: uv add python-docx"
            )

        doc = Document(file_stream)
        paragraphs = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        if not paragraphs:
            raise ExtractionError("No text could be extracted from DOCX")

        full_text = "\n\n".join(paragraphs)
        return ExtractedContent.from_text(full_text)

    def _extract_text_file(self, file_path: Path) -> ExtractedContent:
        """Extract text from a plain text or markdown file."""
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]

        for encoding in encodings:
            try:
                text = file_path.read_text(encoding=encoding)
                if text.strip():
                    return ExtractedContent.from_text(text.strip())
            except UnicodeDecodeError:
                continue

        raise ExtractionError("Could not decode text file with supported encodings")

    def _extract_text_from_stream(self, file_stream: BinaryIO) -> ExtractedContent:
        """Extract text from a text file stream."""
        content = file_stream.read()
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]

        for encoding in encodings:
            try:
                text = content.decode(encoding)
                if text.strip():
                    return ExtractedContent.from_text(text.strip())
            except UnicodeDecodeError:
                continue

        raise ExtractionError("Could not decode text file with supported encodings")
